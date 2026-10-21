#!/usr/bin/env python3
"""Convert bounded external tennis files to the existing Sackmann-style linker input.

SOURCE-ONLY: never writes CDB, resolves a canonical ID, or claims import readiness.
Unknown columns/formats and malformed rows are quarantined, never guessed.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path

MAX_FILE_BYTES = 64 * 1024 * 1024
MAX_MEMBER_BYTES = 32 * 1024 * 1024
MAX_ROWS = 200_000
MAX_MEMBERS = 100
EXTENSIONS = {".csv", ".tsv", ".json", ".jsonl", ".xlsx", ".parquet", ".zip"}
FIELDS = (
    "league", "tourney_date", "tourney_id", "match_num", "tourney_name",
    "winner_name", "loser_name", "surface", "round", "best_of",
    # Retain source identifiers as unverified provenance, never authoritative IDs.
    "source_winner_id", "source_loser_id",
    "w_ace", "w_df", "w_svpt", "w_1stIn", "w_1stWon", "w_2ndWon",
    "w_bpSaved", "w_bpFaced", "l_ace", "l_df", "l_svpt",
    "l_1stIn", "l_1stWon", "l_2ndWon", "l_bpSaved", "l_bpFaced",
)
ALIASES = {
    "league": ("league", "tour", "circuit"),
    "tourney_date": ("tourney_date", "match_date", "event_date", "date"),
    "tourney_id": ("tourney_id", "tournament_id", "event_id"),
    "match_num": ("match_num", "match_id", "match_number"),
    "tourney_name": ("tourney_name", "tournament", "tournament_name", "event"),
    "winner_name": ("winner_name", "winner", "winning_player"),
    "loser_name": ("loser_name", "loser", "losing_player"),
    "surface": ("surface", "court_surface"),
    "round": ("round", "round_name"),
    "best_of": ("best_of", "bestof"),
    "source_winner_id": ("winner_id", "winner_player_id", "winning_player_id"),
    "source_loser_id": ("loser_id", "loser_player_id", "losing_player_id"),
}
STAT_FIELDS = FIELDS[12:]
for field in STAT_FIELDS:
    ALIASES[field] = (field,)
    if field.endswith("_ace"):
        ALIASES[field] = (field, field.replace("_ace", "_aces"))
    if field.endswith("_df"):
        ALIASES[field] = (field, field.replace("_df", "_double_faults"))
def key(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())
def digest(blob):
    return hashlib.sha256(blob).hexdigest()
def iso_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y%m%d")
    raw = str(value or "").strip()
    # Only known unambiguous calendar representations. Do not guess US/UK
    # slash dates or treat numeric Unix timestamps/Excel serials as match dates.
    if re.fullmatch(r"[0-9]{8}\.0+", raw):
        raw = raw.split(".", 1)[0]
    for fmt in ("%Y%m%d", "%Y-%m-%d", "%Y/%m/%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(raw[:10], fmt).strftime("%Y%m%d")
        except ValueError:
            pass
    return ""
def tour_from_filename(name):
    lowered = Path(name).name.lower()
    matches = [x for x in ("atp", "wta") if re.search(r"(^|[^a-z])" + x + r"([^a-z]|$)", lowered)]
    return matches[0] if len(matches) == 1 else ""
def normalize(raw, filename, index):
    if not isinstance(raw, dict):
        return None, "not_object"
    lookup = {key(k): v for k, v in raw.items() if k is not None}
    row = {}
    for target, aliases in ALIASES.items():
        value = next((lookup[key(alias)] for alias in aliases if key(alias) in lookup), "")
        row[target] = "" if value is None else str(value).strip()
    row["tourney_date"] = iso_date(next((lookup[key(a)] for a in ALIASES["tourney_date"] if key(a) in lookup), ""))
    # Some match-level ATP exports expose the exact UTC start as well as a
    # tournament start date. Exact UTC evidence replaces only the date used for
    # linking; never infer a match date from the canonical history itself.
    raw_start = lookup.get(key("start_time_utc"))
    if raw_start is not None and str(raw_start).strip():
        start = str(raw_start).strip()
        if not re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}T[0-9]{2}:[0-9]{2}:[0-9]{2}Z", start):
            return None, "start_time_utc_invalid"
        try:
            observed = datetime.strptime(start, "%Y-%m-%dT%H:%M:%SZ")
            tournament = datetime.strptime(row["tourney_date"], "%Y%m%d")
        except ValueError:
            return None, "start_time_utc_invalid"
        # Dated season containers (e.g. Davis Cup) can point hundreds of
        # days before a real match; do not silently certify such disagreements.
        if not 0 <= (observed.date() - tournament.date()).days <= 16:
            return None, "start_time_outside_tournament_window"
        explicit_match = lookup.get(key("match_date"))
        if explicit_match is not None and str(explicit_match).strip():
            if iso_date(explicit_match) != observed.strftime("%Y%m%d"):
                return None, "match_date_start_time_conflict"
        row["tourney_date"] = observed.strftime("%Y%m%d")
    row["league"] = (row["league"] or tour_from_filename(filename)).lower()
    if row["league"] not in {"atp", "wta"}:
        return None, "tour_missing_or_unsupported"
    if not row["tourney_date"]:
        return None, "date_missing_or_invalid"
    if not row["winner_name"] or not row["loser_name"] or key(row["winner_name"]) == key(row["loser_name"]):
        return None, "players_missing_or_equal"
    if not row["tourney_name"]:
        return None, "tournament_missing"
    if not any(row[f] for f in STAT_FIELDS):
        return None, "no_supported_serve_return_stats"
    row["match_num"] = row["match_num"] or str(index)
    return row, ""
def read_rows(blob, name):
    suffix = Path(name).suffix.lower()
    if suffix in {".csv", ".tsv"}:
        sample = blob.decode("utf-8-sig")
        delimiter = "\t" if suffix == ".tsv" else ","
        return csv.DictReader(io.StringIO(sample, newline=""), delimiter=delimiter)
    if suffix == ".jsonl":
        return [json.loads(line) for line in blob.decode("utf-8-sig").splitlines() if line.strip()]
    if suffix == ".json":
        data = json.loads(blob.decode("utf-8-sig"))
        if isinstance(data, dict):
            data = data.get("matches", data.get("rows"))
        if not isinstance(data, list):
            raise ValueError("JSON must contain a list or matches/rows list")
        return data
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
        try:
            if len(wb.worksheets) != 1:
                raise ValueError("Excel must contain exactly one data sheet")
            it = wb.worksheets[0].iter_rows(values_only=True)
            header = next(it, None)
            if not header:
                return []
            keys = [str(v or "") for v in header]
            if len(set(keys)) != len(keys):
                raise ValueError("Duplicate Excel headers")
            return [dict(zip(keys, row)) for row in it]
        finally:
            wb.close()
    if suffix == ".parquet":
        import pandas as pd
        return pd.read_parquet(io.BytesIO(blob)).to_dict("records")
    raise ValueError("Unsupported table format")
def pieces(name, blob):
    if Path(name).suffix.lower() != ".zip":
        yield name, blob
        return
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        members = [x for x in archive.infolist() if not x.is_dir()]
        if len(members) > MAX_MEMBERS:
            raise ValueError("Archive has too many members")
        if sum(x.file_size for x in members) > MAX_FILE_BYTES:
            raise ValueError("Archive exceeds uncompressed size limit")
        for info in members:
            suffix = Path(info.filename).suffix.lower()
            if suffix not in EXTENSIONS - {".zip"}:
                continue
            if info.file_size > MAX_MEMBER_BYTES or info.flag_bits & 1:
                raise ValueError("Oversized or encrypted ZIP member")
            yield info.filename, archive.read(info)
def run(input_dir: Path, out_dir: Path, max_files: int = 100):
    out_dir.mkdir(parents=True, exist_ok=True)
    normalized = out_dir / "normalized"
    normalized.mkdir(exist_ok=True)
    files = sorted(p for p in input_dir.iterdir() if p.is_file())
    if len(files) > max_files:
        raise ValueError("Unbounded inbox: file count exceeds configured limit")
    counts = Counter()
    review = []
    outputs = []
    source_diagnostics = []
    hashes = set()
    for path in files:
        blob = path.read_bytes()
        sha = digest(blob)
        if sha in hashes:
            counts["duplicate_source_files"] += 1
            continue
        hashes.add(sha)
        if len(blob) > MAX_FILE_BYTES or path.suffix.lower() not in EXTENSIONS:
            counts["unsupported_or_oversized_files"] += 1
            review.append({"sha256": sha, "reason": "unsupported_or_oversized"})
            continue
        try:
            components = list(pieces(path.name, blob))
        except (ValueError, zipfile.BadZipFile, RuntimeError):
            counts["invalid_archives"] += 1
            review.append({"sha256": sha, "reason": "invalid_archive"})
            continue
        for part_no, (member_name, part) in enumerate(components):
            if len(part) > MAX_MEMBER_BYTES:
                counts["unsupported_or_oversized_files"] += 1
                continue
            counts["source_tables"] += 1
            try:
                rows = read_rows(part, member_name)
                tables = {"atp": [], "wta": []}
                table_rejections = Counter()
                table_seen = 0
                for index, raw in enumerate(rows, 1):
                    if index > MAX_ROWS:
                        raise ValueError("Table exceeds row limit")
                    counts["raw_rows"] += 1
                    table_seen += 1
                    row, reason = normalize(raw, member_name, index)
                    if reason:
                        counts["rejected_rows"] += 1
                        counts["reason_" + reason] += 1
                        table_rejections[reason] += 1
                        continue
                    tables[row["league"]].append(row)
                for tour, accepted in tables.items():
                    if not accepted:
                        continue
                    dest = normalized / (tour + "_" + sha[:16] + "_" + str(part_no) + ".csv")
                    with dest.open("w", encoding="utf-8", newline="") as handle:
                        writer = csv.DictWriter(handle, fieldnames=FIELDS)
                        writer.writeheader()
                        writer.writerows(accepted)
                    outputs.append({"path": str(dest), "tour": tour, "rows": len(accepted),
                                    "source_sha256": sha, "member_sha256": digest(part)})
                    counts["normalized_rows"] += len(accepted)
                source_diagnostics.append({
                    "source_sha256": sha, "member_sha256": digest(part),
                    "raw_rows": table_seen,
                    "normalized_rows": table_seen - sum(table_rejections.values()),
                    "rejected_by_reason": dict(sorted(table_rejections.items())),
                    "state": "SOURCE_ONLY", "production_mutated": False,
                })
            except (ValueError, KeyError, UnicodeError, csv.Error, ImportError, TypeError) as exc:
                counts["review_tables"] += 1
                review.append({"sha256": sha, "member_sha256": digest(part),
                               "reason": type(exc).__name__})
    report = {"schema": 1, "state": "SOURCE_ONLY", "production_mutated": False,
              "api_requests": 0, "counts": dict(counts),
              "normalized": outputs, "review": review,
        "source_diagnostics": source_diagnostics}
    (out_dir / "report.json").write_text(json.dumps(report, indent=2, sort_keys=True) + "\n")
    return report
def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    parser.add_argument("--max-files", type=int, default=100)
    args = parser.parse_args()
    report = run(args.input_dir, args.out_dir, args.max_files)
    print(json.dumps({"state": report["state"], "counts": report["counts"],
                      "normalized_files": len(report["normalized"])}))
if __name__ == "__main__":
    main()
